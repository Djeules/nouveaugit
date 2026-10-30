"""Construit la charte de marque de Nine au format PowerPoint (16:9).

Le fichier s'importe dans Canva (Créer un design > Importer) : textes, formes
et couleurs restent modifiables. Les logos sont des images PNG, à remplacer
par les SVG du dossier logos/ si besoin.

    python3 logos.py      # d'abord, pour générer les logos
    python3 charte.py     # produit nine-charte-de-marque.pptx

Unités : la page fait 1920 × 1080 « unités » ; une unité = 0,5 pt.
Les tailles de police sont données dans ces mêmes unités.
"""
import pathlib
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ICI = pathlib.Path(__file__).parent
LOGOS = ICI / "logos"
POLICES = pathlib.Path.home() / ".fonts" / "dm"
SORTIE = ICI / "nine-charte-de-marque.pptx"

# ── Couleurs ───────────────────────────────────────────────────────────────
ENCRE, CREME, CARTON, KRAFT = "1F1C18", "F4EEE4", "E5D3B8", "BF9465"
TOMATE, CRAYON, BLANC = "E5452A", "6B645B", "FFFFFF"

SERIF, SANS = "DM Serif Display", "DM Sans"
MESURE = {  # pour calculer la largeur des étiquettes
    (SANS, True): POLICES / "f4.ttf", (SANS, False): POLICES / "f2.ttf",
    (SERIF, False): POLICES / "f6.ttf",
}

U = 6350                      # EMU par unité
W, H = 1920, 1080
M = 96                        # marge
COL = (W - 2 * M - 2 * 48) / 3   # trois colonnes, gouttière 48
C1, C2, C3 = M, M + COL + 48, M + 2 * (COL + 48)


def fr(t):
    """Espaces insécables de la typographie française."""
    for a, b in ((" :", " :"), (" ;", " ;"), (" ?", " ?"),
                 (" !", " !"), ("« ", "« "), (" »", " »"),
                 (" €", " €"), (" %", " %"), (" h ", " h ")):
        t = t.replace(a, b)
    return t


def largeur(texte, police=SANS, taille=24, gras=True, espacement=0):
    f = ImageFont.truetype(str(MESURE[(police, gras)]), int(taille * 4))
    return f.getlength(texte) / 4 + espacement * (len(texte) - 1)


# ── Primitives ─────────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(W * U), Emu(H * U)
VIDE = prs.slide_layouts[6]
pages = []


def e(v):
    return Emu(int(round(v * U)))


def page(fond=CREME):
    s = prs.slides.add_slide(VIDE)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(fond)
    pages.append(s)
    return s


def sans_style(forme):
    """Retire le style de thème (ombre portée, contour bleu) : Canva l'importerait."""
    st = forme._element.find(qn("p:style"))
    if st is not None:
        forme._element.remove(st)


def rect(s, x, y, w, h, fond=None, trait=None, ep=3, forme=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(forme, e(x), e(y), e(w), e(h))
    sans_style(r)
    if fond:
        r.fill.solid()
        r.fill.fore_color.rgb = RGBColor.from_string(fond)
    else:
        r.fill.background()
    if trait:
        r.line.color.rgb = RGBColor.from_string(trait)
        r.line.width = e(ep)
    else:
        r.line.fill.background()
    return r


def ligne(s, x1, y1, x2, y2, couleur=ENCRE, ep=3):
    l = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, e(x1), e(y1), e(x2), e(y2))
    sans_style(l)
    l.line.color.rgb = RGBColor.from_string(couleur)
    l.line.width = e(ep)
    return l


def texte(s, x, y, w, h, contenu, police=SANS, taille=28, couleur=ENCRE,
          gras=False, align="g", ancre="h", interligne=None, espacement=0,
          apres=0, italique=False):
    """contenu : str, ou liste de paragraphes ; un paragraphe est un str ou
    une liste de segments (texte, {options}) pour mélanger les styles."""
    b = s.shapes.add_textbox(e(x), e(y), e(w), e(h))
    tf = b.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = {"h": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[ancre]
    paras = contenu if isinstance(contenu, list) else [contenu]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"g": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "d": PP_ALIGN.RIGHT}[align]
        if interligne:
            p.line_spacing = interligne
        if apres:
            p.space_after = Pt(apres * 0.5)
        segs = para if isinstance(para, list) else [(para, {})]
        for seg, opt in segs:
            r = p.add_run()
            r.text = fr(seg)
            f = r.font
            f.name = opt.get("police", police)
            f.size = Pt(opt.get("taille", taille) * 0.5)
            f.bold = opt.get("gras", gras)
            f.italic = opt.get("italique", italique)
            f.color.rgb = RGBColor.from_string(opt.get("couleur", couleur))
            esp = opt.get("espacement", espacement)
            if esp:
                r._r.get_or_add_rPr().set("spc", str(int(esp * 50)))
    return b


def image(s, nom, x, y, w=None, h=None):
    return s.shapes.add_picture(str(LOGOS / f"{nom}.png"), e(x), e(y),
                                e(w) if w else None, e(h) if h else None)


def logo(s, nom, x, y, w):
    """Le mot « nine » : PNG de 2400 × 897 px."""
    return image(s, nom, x, y, w=w)


def etiquette(s, x, y, t, fond=ENCRE, coul=CREME, taille=22, ancre_droite=False):
    """La bande d'étiqueteuse : capitales espacées sur une bande pleine."""
    t = fr(t.upper())
    esp = taille * 0.14
    w = largeur(t, SANS, taille, True, esp) + taille * 1.3
    h = taille * 1.9
    if ancre_droite:
        x -= w
    rect(s, x, y, w, h, fond)
    texte(s, x, y, w, h, t, SANS, taille, coul, gras=True, align="c", ancre="m", espacement=esp)
    return w


def chrome(s, chapitre, sombre=False, couleur=None):
    """En-tête et numéro de page, comme une page du site de marque."""
    c = couleur or (CREME if sombre else CRAYON)
    texte(s, M, 44, 700, 30, "nine · charte de marque", SANS, 20, c, gras=True, espacement=2)
    texte(s, W - M - 700, 44, 700, 30, chapitre, SANS, 20, c, gras=True, align="d", espacement=2)
    texte(s, W - M - 100, H - 64, 100, 30, f"{len(pages):02d}", SANS, 20, c, gras=True, align="d")


def titre(s, t, y=120, w=1400, taille=96, couleur=ENCRE, h=None):
    return texte(s, M, y, w, h or taille * 2.4, t, SERIF, taille, couleur, interligne=0.95)


def chapo(s, t, y=260, w=1100, couleur=ENCRE, taille=32):
    return texte(s, M, y, w, 200, t, SANS, taille, couleur, interligne=1.25)


def photo(s, x, y, w, h, legende, fond=CARTON):
    """Emplacement photo : à remplacer dans Canva (glisser l'image dessus)."""
    rect(s, x, y, w, h, fond)
    ligne(s, x, y, x + w, y + h, KRAFT, 1.5)
    ligne(s, x + w, y, x, y + h, KRAFT, 1.5)
    tw = min(w - 40, 520)
    rect(s, x + (w - tw) / 2, y + h / 2 - 50, tw, 100, fond)
    texte(s, x + (w - tw) / 2, y + h / 2 - 50, tw, 100,
          [[("PHOTO", {"gras": True, "espacement": 3, "taille": 18})], legende],
          SANS, 20, ENCRE, align="c", ancre="m")


def ouverture(num, nom, intro, liste):
    """Page d'ouverture de chapitre : fond encre, gros chiffre, le coin tomate."""
    s = page(ENCRE)
    rect(s, W - 280, 0, 280, 280, TOMATE)            # le coin de neuf
    texte(s, M - 12, 120, 900, 700, num, SERIF, 620, CREME, interligne=0.8)
    texte(s, 1000, 300, 820, 290, nom, SERIF, 104, CREME, interligne=0.95, ancre="b")
    texte(s, 1000, 620, 760, 180, intro, SANS, 30, CREME, interligne=1.3)
    par_col = 4 if len(liste) > 4 else len(liste)
    for i, l in enumerate(liste):
        col, rang = divmod(i, par_col)
        texte(s, 1000 + col * 400, 830 + rang * 42, 390, 40, [[(f"{num}.{i + 1}   ", {"couleur": KRAFT}), (l, {})]],
              SANS, 22, CREME, gras=True)
    texte(s, M, H - 64, 700, 30, "nine · charte de marque", SANS, 20, CRAYON, gras=True, espacement=2)
    texte(s, W - M - 100, H - 64, 100, 30, f"{len(pages):02d}", SANS, 20, CRAYON, gras=True, align="d")
    return s


def carte(s, x, y, w, h, surtitre, corps, fond=None, trait=ENCRE, couleur=ENCRE, titre_taille=40,
          titre_police=SERIF, corps_taille=24):
    """Un bloc de contenu : un mot-clé en grand, un texte court dessous."""
    if fond or trait:
        rect(s, x, y, w, h, fond, trait, 2)
    p = 32
    texte(s, x + p, y + p, w - 2 * p, titre_taille * 2.6, surtitre, titre_police, titre_taille, couleur,
          interligne=1.0)
    texte(s, x + p, y + p + titre_taille * 1.35 + 24, w - 2 * p, h - titre_taille * 1.4 - 2 * p - 24,
          corps, SANS, corps_taille, couleur, interligne=1.3)


def croix(s, x, y, t=36, couleur=TOMATE):
    m = t * 0.22
    ligne(s, x + m, y + m, x + t - m, y + t - m, couleur, t * 0.11)
    ligne(s, x + t - m, y + m, x + m, y + t - m, couleur, t * 0.11)


def coche(s, x, y, t=36, couleur=ENCRE):
    ligne(s, x + t * 0.15, y + t * 0.52, x + t * 0.4, y + t * 0.76, couleur, t * 0.11)
    ligne(s, x + t * 0.4, y + t * 0.76, x + t * 0.88, y + t * 0.22, couleur, t * 0.11)


# ══════════════════════════════════════════════════════════════════════════
# 01 · Couverture
# ══════════════════════════════════════════════════════════════════════════
s = page(CREME)
# Le carré en neuf cases, à droite : une case tomate, en haut à droite.
gx, gy, cs, gap = W - M - 3 * 268 - 24, 120, 268, 12
for r in range(3):
    for c in range(3):
        coul = TOMATE if (r, c) == (0, 2) else ENCRE
        rect(s, gx + c * (cs + gap), gy + r * (cs + gap), cs, cs, coul)
logo(s, "nine-logo-encre", M, 110, 560)
texte(s, M, 360, 860, 300, "Charte de marque", SERIF, 136, ENCRE, interligne=0.9)
texte(s, M, 660, 760, 160, "Le logo, les couleurs, la voix et la manière de tout ranger au carré.",
      SANS, 32, ENCRE, interligne=1.3)
etiquette(s, M, 880, "Version 1 · octobre 2026")
texte(s, M, 940, 860, 40, "Un document de travail de Nine et Julien, à valider ensemble.", SANS, 22, CRAYON)

# ══════════════════════════════════════════════════════════════════════════
# 02 · Sommaire
# ══════════════════════════════════════════════════════════════════════════
s = page(CREME)
chrome(s, "Sommaire")
titre(s, "Neuf chapitres.\nForcément.", y=120, w=760, taille=104)
texte(s, M, 400, 640, 300,
      "Nine, c'est un prénom, un chiffre et une promesse de neuf. La charte suit le même plan : "
      "neuf cases, une par chapitre. Celle du coin, en haut à droite, c'est le logo.",
      SANS, 28, ENCRE, interligne=1.35)
chapitres = [
    ("01", "La marque", "Mission, promesse, offres"),
    ("02", "La voix", "Comment parle Nine"),
    ("03", "Le logo", "Le point carré, le symbole"),
    ("04", "La couleur", "Une touche, pas une couche"),
    ("05", "La typographie", "Deux polices, pas trois"),
    ("06", "La grille", "Tout tient dans neuf cases"),
    ("07", "La photo", "Lumière du jour, de face"),
    ("08", "Le trait et la matière", "Pictos, étiquettes, kraft"),
    ("09", "En pratique", "Cartes, flyers, réseaux, site"),
]
gx, gy, cs, gap = 900, 110, 288, 12
for i, (n, nom, sous) in enumerate(chapitres):
    r, c = divmod(i, 3)
    x, y = gx + c * (cs + gap), gy + r * (cs + gap)
    coin = (r, c) == (0, 2)
    rect(s, x, y, cs, cs, TOMATE if coin else None, None if coin else ENCRE, 2)
    texte(s, x + 24, y + 20, cs - 48, 60, n, SERIF, 56, ENCRE)
    texte(s, x + 24, y + cs - 120, cs - 48, 100,
          [[(nom, {"gras": True, "taille": 26})], sous], SANS, 20, ENCRE, interligne=1.2)

# ══════════════════════════════════════════════════════════════════════════
# 03 · Mode d'emploi
# ══════════════════════════════════════════════════════════════════════════
s = page(CREME)
chrome(s, "Avant de commencer")
titre(s, "Une marque se range\ncomme un placard.", w=1300)
chapo(s, "Chaque chose à sa place, et une place pour chaque chose. Cette charte dit où va quoi : "
         "le logo, les couleurs, les mots. Pour que, d'un post Instagram à une carte de visite, "
         "on reconnaisse Nine avant même d'avoir lu son nom.", y=380, w=1150)
for i, (t, c) in enumerate([
    ("Pour Nine", "Les visuels Canva, les réseaux, les légendes, les messages aux clients."),
    ("Pour Julien", "Le site, les flyers, la prospection, les documents (devis, CGV, fiches)."),
    ("Pour les autres", "Imprimeur, brodeur de tenues, partenaires : la page 04 et le chapitre 03 suffisent."),
]):
    x = [C1, C2, C3][i]
    ligne(s, x, 700, x + COL, 700, ENCRE, 3)
    texte(s, x, 730, COL, 60, t, SERIF, 44, ENCRE)
    texte(s, x, 800, COL - 40, 160, c, SANS, 26, ENCRE, interligne=1.35)
texte(s, M, 980, 1500, 40,
      "Dans Canva : chaque texte, forme et couleur de ce document se modifie. Les photos sont des emplacements à remplacer.",
      SANS, 20, CRAYON)

# ══════════════════════════════════════════════════════════════════════════
# 04 · La marque en neuf cases (le cadre, à la manière de Dropbox)
# ══════════════════════════════════════════════════════════════════════════
s = page(CREME)
chrome(s, "La marque en un coup d'œil")
titre(s, "La marque\nen neuf cases.", w=700, taille=96)
texte(s, M, 380, 620, 400,
      "Chaque case est un élément de la marque. Seul, aucun ne suffit. Ensemble, ils font Nine. "
      "Si vous ne deviez garder qu'une page de ce document, ce serait celle-ci.",
      SANS, 28, ENCRE, interligne=1.35)
gx, gy, cs, gap = 836, 110, 286, 12
cs_w = (W - M - gx - 2 * gap) / 3
for i in range(9):
    r, c = divmod(i, 3)
    x, y = gx + c * (cs_w + gap), gy + r * (cs + gap)
    if i == 0:      # voix
        rect(s, x, y, cs_w, cs, ENCRE)
        texte(s, x + 24, y + 24, cs_w - 48, 200, "Vous n'aimez pas ranger. Nine, si.", SERIF, 40, CREME, interligne=1.0)
        texte(s, x + 24, y + cs - 50, cs_w, 30, "LA VOIX", SANS, 18, KRAFT, gras=True, espacement=3)
    elif i == 1:    # typo
        rect(s, x, y, cs_w, cs, BLANC)
        texte(s, x + 24, y + 4, cs_w, 200, "Aa", SERIF, 170, ENCRE)
        texte(s, x + 24, y + cs - 50, cs_w, 30, "LA TYPOGRAPHIE", SANS, 18, CRAYON, gras=True, espacement=3)
    elif i == 2:    # le coin
        rect(s, x, y, cs_w, cs, TOMATE)
        logo(s, "nine-logo-sur-tomate", x + 34, y + 70, cs_w - 68)
        texte(s, x + 24, y + cs - 50, cs_w, 30, "LE LOGO", SANS, 18, ENCRE, gras=True, espacement=3)
    elif i == 3:    # couleur
        bandes = [(CREME, .44), (ENCRE, .26), (CARTON, .12), (KRAFT, .10), (TOMATE, .08)]
        bx = x
        for coul, part in bandes:
            rect(s, bx, y, cs_w * part, cs, coul, None)
            bx += cs_w * part
        rect(s, x, y, cs_w, cs, None, ENCRE, 1.5)
        rect(s, x + 16, y + cs - 58, 170, 40, CREME)
        texte(s, x + 24, y + cs - 50, cs_w, 30, "LA COULEUR", SANS, 18, ENCRE, gras=True, espacement=3)
    elif i == 4:    # grille
        rect(s, x, y, cs_w, cs, CREME, ENCRE, 2)
        for k in (1, 2):
            ligne(s, x + cs_w * k / 3, y, x + cs_w * k / 3, y + cs, ENCRE, 1.5)
            ligne(s, x, y + cs * k / 3, x + cs_w, y + cs * k / 3, ENCRE, 1.5)
        rect(s, x + cs_w * 2 / 3, y, cs_w / 3, cs / 3, ENCRE)
        rect(s, x + 12, y + cs - 58, 140, 40, CREME)
        texte(s, x + 20, y + cs - 50, cs_w, 30, "LA GRILLE", SANS, 18, ENCRE, gras=True, espacement=3)
    elif i == 5:    # photo
        photo(s, x, y, cs_w, cs, "De face, en lumière du jour")
        texte(s, x + 24, y + cs - 50, cs_w, 30, "LA PHOTO", SANS, 18, ENCRE, gras=True, espacement=3)
    elif i == 6:    # étiquette
        rect(s, x, y, cs_w, cs, KRAFT)
        etiquette(s, x + 24, y + 30, "à garder", taille=24)
        etiquette(s, x + 24, y + 86, "à donner", taille=24)
        etiquette(s, x + 24, y + 142, "à jeter", taille=24, fond=CREME, coul=ENCRE)
        texte(s, x + 24, y + cs - 50, cs_w, 30, "LA MATIÈRE", SANS, 18, ENCRE, gras=True, espacement=3)
    elif i == 7:    # trait
        rect(s, x, y, cs_w, cs, BLANC)
        bx, by = x + cs_w / 2 - 70, y + 60
        rect(s, bx - 10, by, 160, 36, None, ENCRE, 4)
        rect(s, bx, by + 36, 140, 100, None, ENCRE, 4)
        rect(s, bx + 45, by + 66, 50, 24, None, ENCRE, 4)
        texte(s, x + 24, y + cs - 50, cs_w, 30, "LE TRAIT", SANS, 18, CRAYON, gras=True, espacement=3)
    elif i == 8:    # mouvement
        rect(s, x, y, cs_w, cs, CARTON)
        for k, a in enumerate((0, 1, 2)):
            rect(s, x + 30 + k * 70, y + 90 + (2 - k) * 22, 56, 56, ENCRE if k < 2 else TOMATE)
        texte(s, x + 24, y + cs - 50, cs_w, 30, "LE MOUVEMENT", SANS, 18, ENCRE, gras=True, espacement=3)

# ══════════════════════════════════════════════════════════════════════════
# CHAPITRE 01 · LA MARQUE
# ══════════════════════════════════════════════════════════════════════════
CH = "01 · La marque"
ouverture("01", "La marque", "Avant les couleurs et le logo : ce que Nine fait, pour qui, et pourquoi personne d'autre ne le fait comme ça.",
          ["Mission, vision, ambition", "La promesse", "Trois mots", "La place de Nine", "Le regard de Nine", "Pour qui", "Les offres"])

# Mission, vision, ambition
s = page(CREME)
chrome(s, CH)
titre(s, "Pourquoi Nine existe.", taille=96)
blocs = [
    ("Mission", "Ranger chez ceux qui n'ont pas envie de le faire. Sans devis, sans leçon, sans qu'ils aient à lever le petit doigt."),
    ("Vision", "Que ranger se délègue comme le ménage ou le repassage : on réserve un créneau, et c'est réglé."),
    ("Ambition", "Devenir le premier réflexe entre Saint-Nazaire et Nantes dès que c'est « un peu le bazar ». Après un premier essai, on revient."),
]
for i, (t, c) in enumerate(blocs):
    x = [C1, C2, C3][i]
    rect(s, x, 330, COL, 560, ENCRE if i == 0 else None, None if i == 0 else ENCRE, 2)
    coul = CREME if i == 0 else ENCRE
    etiquette(s, x + 32, 362, t, fond=TOMATE if i == 0 else ENCRE, coul=ENCRE if i == 0 else CREME)
    texte(s, x + 32, 450, COL - 64, 420, c, SERIF, 46, coul, interligne=1.08)
texte(s, M, 940, 1600, 60,
      "La mission dit ce qu'on fait demain matin. La vision, ce qu'on change. L'ambition, où on veut être dans trois ans.",
      SANS, 24, CRAYON)

# La promesse
s = page(TOMATE)
chrome(s, CH, couleur=ENCRE)
texte(s, M, 110, 1000, 40, "LA PROMESSE", SANS, 22, ENCRE, gras=True, espacement=4)
texte(s, M - 8, 180, 1700, 460, "Vous réservez.\nNine range.", SERIF, 230, ENCRE, interligne=0.88)
for i, (k, t, ex) in enumerate([
    ("1 · Le résultat", "Un coin de chez vous rangé", "Ce que le client obtient."),
    ("2 · Le délai", "en deux heures,", "Combien de temps ça prend."),
    ("3 · Zéro effort", "sans lever le petit doigt.", "Ce qu'il a à faire : rien."),
]):
    x = [C1, C2, C3][i]
    ligne(s, x, 740, x + COL, 740, ENCRE, 3)
    texte(s, x, 760, COL, 30, k.upper(), SANS, 20, ENCRE, gras=True, espacement=3)
    texte(s, x, 805, COL, 110, t, SERIF, 44, ENCRE, interligne=1.0)
    texte(s, x, 940, COL, 60, ex, SANS, 24, ENCRE)

# Trois mots
s = page(CREME)
chrome(s, CH)
titre(s, "Trois mots qui\ntranchent les débats.", w=1300)
chapo(s, "Quand on hésite entre deux visuels, deux phrases ou deux idées, on garde celle qui coche les trois.",
      y=360)
for i, (m, d, oui, non) in enumerate([
    ("Limpide.", "Une offre se comprend en une phrase. Un visuel se lit en une seconde.",
     "« Un coin. 2 heures. 100 €. »", "« Notre approche holistique de l'espace »"),
    ("Carré.", "Tout a sa place : grille, symétrie, alignements. Comme les LEGO de Nine.",
     "Une grille, des angles droits, du vide.", "Des formes qui partent dans tous les sens."),
    ("Discret.", "On entre chez les gens. On ne juge pas, on ne montre pas, on ne raconte pas.",
     "Des mains, des tiroirs, des étagères.", "Le visage ou l'adresse d'un client."),
]):
    x = [C1, C2, C3][i]
    texte(s, x, 500, COL, 100, m, SERIF, 84, TOMATE if i == 0 else ENCRE)
    texte(s, x, 610, COL - 20, 120, d, SANS, 26, ENCRE, interligne=1.35)
    ligne(s, x, 760, x + COL, 760, ENCRE, 2)
    coche(s, x, 785, 28)
    texte(s, x + 44, 785, COL - 44, 70, oui, SANS, 22, ENCRE, interligne=1.3)
    croix(s, x, 870, 28)
    texte(s, x + 44, 870, COL - 44, 70, non, SANS, 22, CRAYON, interligne=1.3)

# La place de Nine
s = page(CREME)
chrome(s, CH)
titre(s, "Ce qu'on vend :\ndu temps, et un coin rangé.", w=1500, taille=88)
chapo(s, "Pas une liste de tâches, pas une « transformation de vie ». Le client réserve comme il réserverait "
         "une femme de ménage, et Nine fait le travail. Elle ne vient pas lui apprendre à ranger : elle range.",
      y=340, w=1500, taille=28)
cols = [
    ("Les home organisers", "Un appel ou une visite, un diagnostic, puis un devis.", "60 à 70 € de l'heure", None),
    ("Nine", "On réserve un créneau en trois clics. Prix affiché, pas de devis.", "50 € de l'heure", TOMATE),
    ("Les plateformes", "Des particuliers non spécialisés, réservés à la tâche.", "dès 12 € de l'heure", None),
]
for i, (qui, comment, prix, fond) in enumerate(cols):
    x = [C1, C2, C3][i]
    y0 = 520 if fond else 560
    rect(s, x, y0, COL, 1000 - y0 - 40, fond, None if fond else ENCRE, 2)
    texte(s, x + 32, y0 + 32, COL - 64, 60, qui, SERIF, 48, ENCRE)
    texte(s, x + 32, y0 + 110, COL - 64, 140, comment, SANS, 26, ENCRE, interligne=1.35)
    texte(s, x + 32, 880, COL - 64, 50, prix, SANS, 26, ENCRE, gras=True)
texte(s, M, 1000, 1600, 40, "Prix publiés relevés le 30 octobre 2026 autour de Nantes et Saint-Nazaire (voir le business plan).",
      SANS, 18, CRAYON)

# Le regard de Nine
s = page(ENCRE)
chrome(s, CH, sombre=True)
texte(s, M, 110, 1000, 40, "CE QUI NE SE COPIE PAS", SANS, 22, KRAFT, gras=True, espacement=4)
texte(s, M - 6, 180, 1150, 520, "« Face à une pièce, la plupart des gens voient du désordre. Moi, je vois un plan. »",
      SERIF, 96, CREME, interligne=1.0)
texte(s, M, 740, 600, 40, "— Nine", SANS, 28, CREME, gras=True)
photo(s, 1360, 140, 464, 600, "Nine, de face, chez elle", fond=CRAYON)
texte(s, M, 830, 1050, 200,
      "Nine a 22 ans, un master de lettres modernes et deux classes de collège derrière elle. Elle a grandi en "
      "rangeant les placards des autres sans qu'on le lui demande. Julien, son père, s'occupe du reste. "
      "Leurs LEGO étaient carrés et symétriques ; ça n'a pas changé.",
      SANS, 26, CREME, interligne=1.4)
texte(s, 1360, 780, 464, 200,
      "La méthode et la rigueur s'apprennent. Ce regard-là, non : c'est lui que la marque doit montrer.",
      SANS, 24, KRAFT, interligne=1.35)

# Pour qui
s = page(CREME)
chrome(s, CH)
titre(s, "Pour tous ceux qui n'aiment\npas ranger. Et qui l'assument.", w=1700, taille=88)
chapo(s, "On ne cible pas des catégories de gens, mais des situations. Trois d'entre elles font vivre l'activité.",
      y=340, w=1500, taille=28)
for i, (t, d, q) in enumerate([
    ("Les parents débordés", "Le cœur de L'habitude : un créneau chaque mois, et on n'y pense plus.", "Les clients réguliers"),
    ("Les changements de vie", "Déménagement, séparation, naissance. Ponctuel, souvent urgent.", "Les gros paniers"),
    ("Les seniors", "Un logement plus simple à vivre. Souvent réservé par leurs enfants.", "Ceux qu'on rassure"),
]):
    x = [C1, C2, C3][i]
    rect(s, x, 480, COL, 460, CARTON if i else ENCRE)
    coul = ENCRE if i else CREME
    texte(s, x + 32, 512, COL - 64, 130, t, SERIF, 52, coul, interligne=1.0)
    texte(s, x + 32, 670, COL - 64, 160, d, SANS, 26, coul, interligne=1.35)
    etiquette(s, x + 32, 870, q, fond=TOMATE if i == 0 else ENCRE, coul=ENCRE if i == 0 else CREME, taille=20)
texte(s, M, 980, 1700, 40, "Le lendemain de soirée, c'est très drôle sur les réseaux. Mais c'est du ménage : pas une cible.",
      SANS, 20, CRAYON)

# Les offres
s = page(CREME)
chrome(s, CH)
titre(s, "Deux offres, pas plus.", taille=96)
chapo(s, "Plus le choix est grand, moins la décision se prend. Les offres se nomment par leur taille, "
         "jamais par une métaphore : on commande « Un coin » sans aucune gêne.", y=250, w=1500, taille=28)
for i, (nom, duree, prix, pour) in enumerate([
    ("Un coin", "2 heures", "100 €", "Un placard, une salle de bain, une entrée, un bureau, une penderie."),
    ("Une pièce", "Une demi-journée", "200 €", "Une cuisine, une chambre, une chambre d'enfant, un dressing, un cellier."),
]):
    x = M + i * 600
    rect(s, x, 400, 568, 520, ENCRE if i == 0 else None, None if i == 0 else ENCRE, 2)
    coul = CREME if i == 0 else ENCRE
    if i == 0:
        rect(s, x + 568 - 80, 400, 80, 80, TOMATE)
    texte(s, x + 40, 440, 480, 100, nom, SERIF, 88, coul)
    texte(s, x + 40, 560, 480, 60, [[(duree, {"gras": True}), ("  ·  " + prix, {})]], SANS, 34, coul)
    texte(s, x + 40, 640, 480, 160, pour, SANS, 26, coul, interligne=1.35)
    texte(s, x + 40, 840, 480, 40, "RÉSERVER", SANS, 22, TOMATE if i == 0 else ENCRE, gras=True, espacement=4)
texte(s, 1340, 400, 484, 40, "AU SECOND PLAN", SANS, 20, CRAYON, gras=True, espacement=3)
for k, (nom, d) in enumerate([
    ("Toute la maison", "Sur devis, à partir de 600 €. Plusieurs pièces, un déménagement."),
    ("L'habitude", "2 heures par mois, le même jour. 100 € par mois, sans engagement."),
]):
    y = 460 + k * 230
    ligne(s, 1340, y, 1824, y, ENCRE, 2)
    texte(s, 1340, y + 20, 484, 60, nom, SERIF, 46, ENCRE)
    texte(s, 1340, y + 90, 484, 120, d, SANS, 24, ENCRE, interligne=1.35)
rect(s, M, 940, 1600, 70, CARTON)
texte(s, M + 24, 940, 1560, 70,
      [[("Le prix après crédit d'impôt (50 € pour Un coin) ", {"gras": True}),
        ("ne s'affiche qu'une fois le récépissé de déclaration services à la personne reçu.", {})]],
      SANS, 22, ENCRE, ancre="m")

# ══════════════════════════════════════════════════════════════════════════
# CHAPITRE 02 · LA VOIX
# ══════════════════════════════════════════════════════════════════════════
CH = "02 · La voix"
ouverture("02", "La voix", "La meilleure chose du site actuel, c'est la façon dont Nine parle. On la garde, et on la descend d'un étage : moins de Flaubert, plus de cuisine.",
          ["Comment parle Nine", "On dit, on évite", "Les phrases qu'on garde", "Dans la vraie vie"])

s = page(CREME)
chrome(s, CH)
titre(s, "Comment parle Nine.", taille=96)
for i, (t, d, ex) in enumerate([
    ("À la première personne.", "C'est Nine qui parle, pas « l'équipe ». Un « je », un prénom, une vraie personne.",
     "« Je range, vous vous asseyez. »"),
    ("La flemme, assumée.", "Personne ne parle comme ça du rangement. On ne culpabilise pas le client : on le comprend.",
     "« Vous n'aimez pas ranger. Nine, si. »"),
    ("Pince-sans-rire.", "De l'humour à froid, à la Jonas Jonasson. Jamais de point d'exclamation pour faire rire.",
     "« Le placard ne va pas se ranger tout seul. On a vérifié. »"),
    ("Des mots de tous les jours.", "On réserve comme on appelle une femme de ménage, pas comme on commande un audit.",
     "« Un créneau, un coin, c'est réglé. »"),
]):
    r, c = divmod(i, 2)
    x, y = M + c * 880, 290 + r * 360
    texte(s, x, y, 820, 70, t, SERIF, 56, TOMATE if i == 1 else ENCRE)
    texte(s, x, y + 90, 780, 120, d, SANS, 26, ENCRE, interligne=1.35)
    texte(s, x, y + 210, 780, 60, ex, SANS, 26, CRAYON, italique=True)

s = page(CREME)
chrome(s, CH)
titre(s, "On dit, on évite.", taille=96)
rect(s, M, 280, 836, 560, ENCRE)
texte(s, M + 40, 310, 600, 40, "ON DIT", SANS, 22, KRAFT, gras=True, espacement=4)
for k, mot in enumerate(["ranger", "réserver", "un créneau", "un coin, une pièce", "Nine"]):
    texte(s, M + 40, 370 + k * 84, 760, 80, mot, SERIF, 60, CREME)
rect(s, 988, 280, 836, 560, None, ENCRE, 2)
texte(s, 1028, 310, 600, 40, "ON ÉVITE", SANS, 22, CRAYON, gras=True, espacement=4)
for k, mot in enumerate(["coach", "home organiser", "désencombrement", "transformation", "accompagnement, bien-être"]):
    texte(s, 1028, 370 + k * 84, 760, 80, mot, SERIF, 60, CRAYON)
    ligne(s, 1028, 370 + k * 84 + 46, 1028 + largeur(mot, SERIF, 60, False), 370 + k * 84 + 46, TOMATE, 4)
texte(s, M, 880, 1728, 140,
      [[("La règle d'or : ", {"gras": True}),
        ("« On range, on ne nettoie pas. » À écrire partout où le client pourrait confondre avec le ménage. "
         "Le coup de chiffon sur ce qu'on vient de vider, lui, est compris.", {})]],
      SANS, 28, ENCRE, interligne=1.35)

s = page(ENCRE)
chrome(s, CH, sombre=True)
texte(s, M, 110, 1000, 40, "LES PHRASES QU'ON GARDE", SANS, 22, KRAFT, gras=True, espacement=4)
phrases = ["Vous n'aimez pas ranger. Nine, si.", "On range. Vous respirez.", "Rangez moins. Vivez plus.",
           "Ranger, ça se délègue."]
for i, p in enumerate(phrases):
    texte(s, M, 200 + i * 180, 1728, 170, p, SERIF, 104 if i == 0 else 92, TOMATE if i == 0 else CREME, interligne=0.95)
texte(s, M, 960, 1728, 60,
      "À saupoudrer sur le site, les réseaux, les flyers. Une phrase par support, jamais deux côte à côte.",
      SANS, 24, KRAFT)

s = page(CREME)
chrome(s, CH)
titre(s, "Dans la vraie vie.", taille=96)
chapo(s, "Même voix partout, du site au SMS. Exemples de ton : à adapter, pas à copier.", y=250, taille=28)
cas = [
    ("Confirmation de réservation", "Votre réservation a bien été enregistrée. Un diagnostic préalable sera effectué.",
     "C'est noté : mardi 14 h, votre entrée. Envoyez-moi trois photos ici, je m'occupe du reste. — Nine"),
    ("La limite du ménage", "Les prestations de nettoyage ne sont pas incluses dans nos offres.",
     "Je range, je ne nettoie pas. Le coup de chiffon sur l'étagère que je viens de vider, en revanche, c'est cadeau."),
    ("Un refus", "Votre demande ne correspond pas à nos critères d'intervention.",
     "Là, il faudrait plus qu'une Nine. Je préfère vous le dire maintenant que vous facturer un après-midi pour un résultat à moitié."),
]
for i, (sit, avant, apres) in enumerate(cas):
    y = 350 + i * 220
    texte(s, M, y, 420, 80, sit, SERIF, 40, ENCRE, interligne=1.0)
    croix(s, 560, y + 4, 26, CRAYON)
    texte(s, 600, y, 520, 180, avant, SANS, 23, CRAYON, interligne=1.35)
    coche(s, 1160, y + 4, 26, TOMATE)
    texte(s, 1200, y, 624, 180, apres, SANS, 23, ENCRE, interligne=1.35, gras=True)
    if i < 2:
        ligne(s, M, y + 200, W - M, y + 200, CARTON, 2)

# ══════════════════════════════════════════════════════════════════════════
# CHAPITRE 03 · LE LOGO
# ══════════════════════════════════════════════════════════════════════════
CH = "03 · Le logo"
ouverture("03", "Le logo", "Un mot en minuscules et un point devenu carré. Le coin de neuf, c'est lui.",
          ["Le logo", "Le symbole", "Les versions", "L'espace et la taille", "Ce qu'on ne fait pas"])

s = page(CREME)
chrome(s, CH)
image(s, "nine-signature-encre", (W - 1060) / 2, 150, w=1060)
texte(s, M, 830, 800, 200,
      [[("Le mot. ", {"gras": True}),
        ("« nine » en minuscules, dans la police des titres, légèrement resserré. Accessible, pas guindé.", {})]],
      SANS, 26, ENCRE, interligne=1.35)
texte(s, 1000, 830, 824, 200,
      [[("Le point. ", {"gras": True, "couleur": TOMATE}),
        ("Le point du i devient un carré tomate : un coin rangé, la seule touche de couleur. "
         "C'est lui qu'on doit retenir.", {})]],
      SANS, 26, ENCRE, interligne=1.35)

s = page(CREME)
chrome(s, CH)
titre(s, "Le symbole.", taille=96)
chapo(s, "Un carré de neuf cases dont une change de couleur : Nine, le chiffre neuf, et un coin remis à neuf. "
         "Il sert là où le mot ne tient pas : photo de profil, icône, tampon, étiquette.", y=250, w=820, taille=28)
# Construction
gx, gy, sz = 1100, 160, 720
rect(s, gx, gy, sz, sz, ENCRE)
rect(s, gx + sz * 2 / 3, gy, sz / 3, sz / 3, TOMATE)
for k in (1, 2):
    ligne(s, gx + sz * k / 3, gy, gx + sz * k / 3, gy + sz, CRAYON, 1.5)
    ligne(s, gx, gy + sz * k / 3, gx + sz, gy + sz * k / 3, CRAYON, 1.5)
texte(s, gx, gy + sz + 20, sz, 40, "1 case = 1/3 du côté. Angles droits, jamais arrondis.", SANS, 22, CRAYON, align="c")
for i, (nom, x) in enumerate([("nine-symbole", M), ("nine-symbole-creme", M + 280), ("nine-symbole-sur-tomate", M + 560)]):
    if i == 1:
        rect(s, x - 20, 560, 240, 240, ENCRE)
    if i == 2:
        rect(s, x - 20, 560, 240, 240, TOMATE)
    image(s, nom, x + 10, 590, 180, 180)
texte(s, M, 830, 820, 120, "Sur crème, sur encre, sur tomate : la case du coin passe toujours à la couleur qui manque.",
      SANS, 24, ENCRE, interligne=1.35)

s = page(CREME)
chrome(s, CH)
titre(s, "Les versions.", taille=96)
cases = [
    (CREME, "nine-logo-encre", "Principale", "Sur crème ou blanc. Dans 9 cas sur 10."),
    (ENCRE, "nine-logo-creme", "Inversée", "Sur encre ou sur photo sombre."),
    (TOMATE, "nine-logo-sur-tomate", "Sur tomate", "Le point passe en crème."),
    (KRAFT, "nine-logo-noir", "Monochrome", "Tampon, gravure, une seule encre."),
]
for i, (fond, nom, t, d) in enumerate(cases):
    x = M + i * (1728 + 32) / 4
    w = (1728 - 3 * 32) / 4
    rect(s, x, 280, w, 440, fond, ENCRE if fond == CREME else None, 2)
    logo(s, nom, x + 40, 400, w - 80)
    texte(s, x, 750, w, 50, t, SERIF, 44, ENCRE)
    texte(s, x, 815, w, 100, d, SANS, 24, ENCRE, interligne=1.3)
texte(s, M, 970, 1728, 40,
      "Les fichiers SVG et PNG de chaque version sont dans le dossier logos/, à importer dans Canva (Importer > fichiers).",
      SANS, 20, CRAYON)

s = page(CREME)
chrome(s, CH)
titre(s, "De l'air autour,\net jamais trop petit.", w=900, taille=88)
# zone de protection
lx, ly, lw = 1040, 200, 640
lh = lw * 897 / 2400
x_unit = lw * 136 / 1907          # côté du carré du point (mot : 1907 unités)
pad = 2 * x_unit
rect(s, lx - pad, ly - pad, lw + 2 * pad, lh + 2 * pad, CARTON)
rect(s, lx - pad, ly - pad, lw + 2 * pad, lh + 2 * pad, None, KRAFT, 2)
logo(s, "nine-logo-encre", lx, ly, lw)
for (xx, yy) in ((lx - pad, ly - pad), (lx + lw, ly - pad), (lx - pad, ly + lh), (lx + lw, ly + lh)):
    rect(s, xx + pad / 4, yy + pad / 4, pad / 2, pad / 2, TOMATE)
texte(s, lx - pad, ly + lh + pad + 20, lw + 2 * pad, 60,
      "Marge minimale : deux fois le carré du point, de chaque côté.", SANS, 22, ENCRE, align="c")
texte(s, M, 420, 820, 200,
      "Le logo respire. Rien ne s'approche à moins de deux carrés de point : ni texte, ni bord de page, ni autre logo.",
      SANS, 28, ENCRE, interligne=1.35)
ligne(s, M, 700, W - M, 700, ENCRE, 2)
for i, (t, d, w) in enumerate([
    ("Logo · impression", "25 mm de large minimum", 250),
    ("Logo · écran", "120 px de large minimum", 190),
    ("Symbole · impression", "8 mm de côté minimum", 80),
    ("Symbole · écran", "32 px de côté minimum", 64),
]):
    x = M + i * 440
    if "Logo" in t:
        logo(s, "nine-logo-encre", x, 740, w * 0.6)
    else:
        image(s, "nine-symbole", x, 740, w * 0.8, w * 0.8)
    texte(s, x, 880, 420, 40, t, SANS, 22, ENCRE, gras=True)
    texte(s, x, 920, 420, 40, d, SANS, 22, ENCRE)

s = page(CREME)
chrome(s, CH)
titre(s, "Ce qu'on ne fait pas.", taille=96)
interdits = [
    "Étirer ou écraser", "Changer la couleur du point", "Arrondir le carré",
    "Faire pivoter", "Ajouter une ombre", "Poser sur une photo chargée",
]
for i, t in enumerate(interdits):
    r, c = divmod(i, 3)
    x, y = [C1, C2, C3][c], 270 + r * 370
    rect(s, x, y, COL, 280, BLANC)
    cx, cy = x + 60, y + 70
    if i == 0:
        image(s, "nine-logo-encre", cx, cy + 20, COL - 120, 120)
    elif i == 1:
        logo(s, "nine-logo-encre", cx + 40, cy, COL - 200)
        k_ = (COL - 200) / 2400       # le point : x 802, y 2,5, côté 171 px sur 2400
        rect(s, cx + 40 + 802 * k_, cy + 2.5 * k_, 171 * k_, 171 * k_, "3A7BD5")
    elif i == 2:
        r_ = rect(s, cx + 150, cy - 20, 170, 170, ENCRE, forme=MSO_SHAPE.ROUNDED_RECTANGLE)
        r_.adjustments[0] = 0.3
        rect(s, cx + 150 + 113, cy - 20, 57, 57, TOMATE, forme=MSO_SHAPE.OVAL)
    elif i == 3:
        p_ = logo(s, "nine-logo-encre", cx + 40, cy, COL - 200)
        p_.rotation = -12
    elif i == 4:
        rect(s, x, y, COL, 280, CARTON)
        logo(s, "nine-logo-noir", cx + 52, cy + 12, COL - 200)
        logo(s, "nine-logo-creme", cx + 40, cy, COL - 200)
    elif i == 5:
        for k in range(6):
            rect(s, x + k * COL / 6, y, COL / 6, 280, [KRAFT, CRAYON, CARTON, TOMATE, ENCRE, KRAFT][k])
        logo(s, "nine-logo-creme", cx + 40, cy + 20, COL - 200)
    croix(s, x, y + 300, 30)
    texte(s, x + 46, y + 300, COL - 46, 40, t, SANS, 24, ENCRE, gras=True)

# ══════════════════════════════════════════════════════════════════════════
# CHAPITRE 04 · LA COULEUR
# ══════════════════════════════════════════════════════════════════════════
CH = "04 · La couleur"
ouverture("04", "La couleur", "Le kraft des boîtes, le crème des murs, l'encre de l'étiqueteuse. Et une touche de tomate, une seule.",
          ["La palette", "Les proportions", "Les bons mélanges"])

s = page(CREME)
chrome(s, CH)
titre(s, "La palette.", taille=96)
palette = [
    ("Tomate", TOMATE, "229 · 69 · 42", "0 · 85 · 90 · 0", "La touche. Le point du logo, un mot, un bouton. Jamais plus d'un élément par visuel."),
    ("Encre", ENCRE, "31 · 28 · 24", "0 · 10 · 20 · 95", "Le texte, les traits, les fonds forts. Un noir chaud, jamais du noir pur."),
    ("Crème", CREME, "244 · 238 · 228", "0 · 3 · 7 · 4", "Le fond. Le mur blanc d'une maison habitée."),
    ("Carton", CARTON, "229 · 211 · 184", "0 · 8 · 22 · 10", "Les blocs secondaires, les encadrés."),
    ("Kraft", KRAFT, "191 · 148 · 101", "0 · 25 · 50 · 25", "La matière : boîtes, étiquettes, papier d'emballage."),
    ("Crayon", CRAYON, "107 · 100 · 91", "0 · 7 · 15 · 58", "Le texte secondaire : légendes, dates, mentions."),
]
w = (1728 - 5 * 16) / 6
for i, (nom, hx, rgb, cmjn, usage) in enumerate(palette):
    x = M + i * (w + 16)
    rect(s, x, 270, w, 400, hx, ENCRE if hx == CREME else None, 2)
    clair = hx in (CREME, CARTON, KRAFT, TOMATE)
    texte(s, x + 20, 290, w - 40, 60, nom, SERIF, 44, ENCRE if clair else CREME)
    texte(s, x + 20, 600, w - 40, 50, f"#{hx}", SANS, 24, ENCRE if clair else CREME, gras=True)
    texte(s, x, 700, w, 120, [[("RVB  ", {"gras": True}), (rgb, {})], [("CMJN  ", {"gras": True}), (cmjn, {})]],
          SANS, 19, ENCRE, interligne=1.4)
    texte(s, x, 790, w - 10, 200, usage, SANS, 20, ENCRE, interligne=1.35)
texte(s, M, 1000, 1728, 30, "Valeurs CMJN indicatives : demander une épreuve à l'imprimeur avant le premier tirage.",
      SANS, 18, CRAYON)

s = page(CREME)
chrome(s, CH)
titre(s, "Une touche,\npas une couche.", w=900, taille=96)
chapo(s, "Sur n'importe quel visuel, la tomate reste minoritaire. C'est ce qui la fait voir. "
         "Si tout est tomate, plus rien ne l'est.", y=380, w=760, taille=28)
parts = [(CREME, 55, "Crème"), (ENCRE, 25, "Encre"), (CARTON, 8, "Carton"), (KRAFT, 7, "Kraft"), (TOMATE, 5, "Tomate")]
bx, by, bw, bh = M, 640, 1728, 220
x = bx
for coul, pct, nom in parts:
    ww = bw * pct / 100
    rect(s, x, by, ww, bh, coul, ENCRE if coul == CREME else None, 2)
    texte(s, x + 12, by + bh + 16, max(ww, 140), 40, f"{nom} · {pct} %", SANS, 22, ENCRE, gras=True)
    x += ww
texte(s, 1000, 200, 824, 400,
      [[("La règle des neuf cases. ", {"gras": True}),
        ("Découpez le visuel en 3 × 3. La tomate n'occupe jamais plus d'une case : "
         "de préférence celle du coin, en haut à droite.", {})]],
      SANS, 28, ENCRE, interligne=1.4)
texte(s, M, 960, 1728, 40, "Le kraft vient d'abord des vraies matières photographiées (boîtes, papier), plus que des aplats.",
      SANS, 20, CRAYON)

s = page(CREME)
chrome(s, CH)
titre(s, "Les bons mélanges.", taille=96)
chapo(s, "Le contraste se mesure, il ne se devine pas. Ratios calculés selon la norme WCAG : 4,5 pour du texte courant, 3 pour un gros titre.",
      y=250, w=1500, taille=26)
paires = [
    (CREME, ENCRE, "Encre sur crème", "14,7", "Tout"),
    (ENCRE, CREME, "Crème sur encre", "14,7", "Tout"),
    (CARTON, ENCRE, "Encre sur carton", "11,6", "Tout"),
    (KRAFT, ENCRE, "Encre sur kraft", "6,2", "Tout"),
    (TOMATE, ENCRE, "Encre sur tomate", "4,2", "Titres seulement"),
    (CREME, TOMATE, "Tomate sur crème", "3,5", "Titres seulement"),
    (CREME, CRAYON, "Crayon sur crème", "5,1", "Légendes"),
    (KRAFT, CREME, "Crème sur kraft", "2,4", "Jamais"),
]
w = (1728 - 3 * 24) / 4
for i, (fond, enc, nom, ratio, usage) in enumerate(paires):
    r, c = divmod(i, 4)
    x, y = M + c * (w + 24), 380 + r * 320
    rect(s, x, y, w, 220, fond, ENCRE if fond == CREME else None, 2)
    texte(s, x + 24, y + 20, w - 48, 120, "Rangez moins.", SERIF, 52, enc)
    texte(s, x + 24, y + 150, w - 48, 50, ratio, SANS, 30, enc, gras=True)
    texte(s, x, y + 236, w - 60, 40, nom, SANS, 22, ENCRE, gras=True)
    texte(s, x, y + 268, w - 60, 40, usage, SANS, 22, TOMATE if usage == "Jamais" else ENCRE)

# ══════════════════════════════════════════════════════════════════════════
# CHAPITRE 05 · LA TYPOGRAPHIE
# ══════════════════════════════════════════════════════════════════════════
CH = "05 · La typographie"
ouverture("05", "La typographie", "Une police à empattements pour parler fort, une police sans pour expliquer. Deux, pas trois.",
          ["Les deux polices", "La hiérarchie", "Les règles d'écriture"])

s = page(CREME)
chrome(s, CH)
titre(s, "Les deux polices.", taille=96)
for i, (pol, role, d) in enumerate([
    (SERIF, "Les titres", "DM Serif Display. Des empattements francs, un peu ronds : sérieux sans être guindé. "
                         "Toujours en un seul corps, jamais en gras ni en capitales."),
    (SANS, "Le texte", "DM Sans. Lisible à toutes les tailles, sur écran comme sur papier. Normal pour le texte, "
                       "gras pour les mots importants et les étiquettes."),
]):
    x = M + i * 880
    rect(s, x, 270, 848, 440, ENCRE if i == 0 else BLANC)
    coul = CREME if i == 0 else ENCRE
    texte(s, x + 40, 260, 600, 340, "Aa", pol, 300, coul, gras=(i == 1))
    texte(s, x + 40, 600, 780, 60, pol, pol, 44, coul, gras=(i == 1))
    texte(s, x, 740, 360, 50, role, SERIF, 44, ENCRE)
    texte(s, x, 810, 820, 160, d, SANS, 24, ENCRE, interligne=1.4)
texte(s, M, 990, 1728, 40,
      "Les deux sont gratuites (Google Fonts), disponibles dans Canva, et s'auto-hébergent sur le site : aucune dépendance.",
      SANS, 20, CRAYON)

s = page(CREME)
chrome(s, CH)
titre(s, "La hiérarchie.", taille=96)
niveaux = [
    ("Titre XL", SERIF, 120, False, 0, "Rangez moins.", "Post 1080 : 120 px · A5 : 48 pt"),
    ("Titre", SERIF, 72, False, 0, "Un coin, en deux heures.", "Post : 72 px · A5 : 28 pt"),
    ("Chapô", SANS, 34, False, 0, "Je viens, je range, vous n'avez rien à faire.", "Post : 34 px · A5 : 13 pt"),
    ("Texte", SANS, 26, False, 0, "Tri, rangement, étiquetage. Sacs « à donner » et « à jeter » préparés.", "Post : 26 px · A5 : 10 pt"),
    ("Étiquette", SANS, 22, True, 4, "RÉSERVER UN CRÉNEAU", "Capitales, gras, espacées"),
]
y = 260
for nom, pol, t, g, esp, ex, regle in niveaux:
    texte(s, M, y + 6, 260, 40, nom.upper(), SANS, 20, CRAYON, gras=True, espacement=3)
    texte(s, 400, y, 1000, t * 1.3, ex, pol, t, ENCRE, gras=g, espacement=esp)
    texte(s, 1424, y + 6, 400, 40, regle, SANS, 20, CRAYON, align="d")
    y += t * 1.3 + 40
    ligne(s, M, y - 20, W - M, y - 20, CARTON, 2)

s = page(CREME)
chrome(s, CH)
titre(s, "Les règles d'écriture.", taille=96)
regles = [
    ("Un titre est une phrase.", "Avec un point final. « Vous réservez. Nine range. » se lit comme une décision, pas comme un slogan."),
    ("Court, toujours.", "Un titre tient en deux lignes. Si ça déborde, on coupe le texte, pas la taille."),
    ("Aligné à gauche.", "Le centré est réservé aux étiquettes et au logo. Le justifié, jamais."),
    ("La typographie française.", "Espace avant : ; ? ! — guillemets « » — 100 € avec une espace."),
    ("Les chiffres en chiffres.", "2 heures, 100 €, 48 h. Plus vite lu, plus vite compris."),
    ("Pas de capitales pour crier.", "Les capitales sont pour les étiquettes. Un titre en capitales, c'est un client qui hausse le ton."),
]
for i, (t, d) in enumerate(regles):
    r, c = divmod(i, 3)
    x, y = [C1, C2, C3][c], 280 + r * 360
    texte(s, x, y, 60, 60, f"{i + 1}", SERIF, 56, TOMATE if i == 0 else KRAFT)
    texte(s, x, y + 80, COL - 20, 60, t, SERIF, 40, ENCRE, interligne=1.0)
    texte(s, x, y + 150, COL - 30, 180, d, SANS, 24, ENCRE, interligne=1.4)

# ══════════════════════════════════════════════════════════════════════════
# CHAPITRE 06 · LA GRILLE
# ══════════════════════════════════════════════════════════════════════════
CH = "06 · La grille"
ouverture("06", "La grille", "Les LEGO de Nine étaient symétriques. Ses visuels aussi : tout se pose dans une grille de neuf cases.",
          ["Neuf cases", "Trois mises en page"])

s = page(CREME)
chrome(s, CH)
titre(s, "Tout tient\ndans neuf cases.", w=820, taille=96)
for i, (t, d) in enumerate([
    ("Trois par trois.", "Chaque format se découpe en 3 × 3, avec une marge de 1/12 du côté."),
    ("Des cases entières.", "Un texte ou une photo occupe 1, 2, 3, 6 ou 9 cases. Jamais une demi-case."),
    ("Le coin est réservé.", "La case en haut à droite accueille la touche : logo, symbole, prix ou étiquette."),
    ("Des angles droits.", "Aucun arrondi, nulle part. Les coins restent des coins."),
]):
    y = 380 + i * 155
    texte(s, M, y, 760, 50, t, SERIF, 42, ENCRE)
    texte(s, M, y + 56, 760, 80, d, SANS, 24, ENCRE, interligne=1.35)
# carré 1080
gx, gy, sz = 1000, 150, 820
rect(s, gx, gy, sz, sz, BLANC, ENCRE, 2)
mg = sz / 12
inner = sz - 2 * mg
rect(s, gx + mg, gy + mg, inner, inner, None, KRAFT, 1.5)
for k in (1, 2):
    ligne(s, gx + mg + inner * k / 3, gy + mg, gx + mg + inner * k / 3, gy + sz - mg, KRAFT, 1.5)
    ligne(s, gx + mg, gy + mg + inner * k / 3, gx + sz - mg, gy + mg + inner * k / 3, KRAFT, 1.5)
rect(s, gx + mg + inner * 2 / 3, gy + mg, inner / 3, inner / 3, TOMATE)
texte(s, gx + mg + inner * 2 / 3, gy + mg, inner / 3, inner / 3, "LE COIN", SANS, 20, ENCRE, gras=True,
      align="c", ancre="m", espacement=3)
texte(s, gx, gy + sz + 16, sz, 40, "Post Instagram 1080 × 1080 · marge 90 px · case 300 px", SANS, 20, CRAYON, align="c")

s = page(CREME)
chrome(s, CH)
titre(s, "Trois mises en page,\net c'est tout.", w=1400, taille=88)
sz = 480
for i, (nom, d) in enumerate([
    ("La phrase", "Une phrase de la voix, sur six cases. Le symbole dans le coin."),
    ("L'avant / après", "Deux photos, même cadrage, chacune sur une demi-largeur. Les étiquettes en bas."),
    ("L'offre", "Le nom, la durée, le prix. Une case tomate, pas deux."),
]):
    x = [C1, C2, C3][i] + (COL - sz) / 2
    y = 360
    rect(s, x, y, sz, sz, BLANC, ENCRE, 2)
    mg = sz / 12
    c3 = (sz - 2 * mg) / 3
    if i == 0:
        rect(s, x, y, sz, sz, ENCRE)
        texte(s, x + mg, y + mg + c3, c3 * 3, c3 * 2, "On range.\nVous respirez.", SERIF, 62, CREME, interligne=0.95)
        image(s, "nine-symbole-creme", x + mg + 2 * c3 + c3 - 60, y + mg, 60, 60)
    elif i == 1:
        photo(s, x, y, sz / 2 - 2, sz, "Avant")
        photo(s, x + sz / 2 + 2, y, sz / 2 - 2, sz, "Après")
        etiquette(s, x + mg, y + sz - mg - 36, "avant", taille=18)
        etiquette(s, x + sz / 2 + mg, y + sz - mg - 36, "après", taille=18, fond=TOMATE, coul=ENCRE)
    else:
        rect(s, x + mg + 2 * c3, y + mg, c3, c3, TOMATE)
        texte(s, x + mg + 2 * c3, y + mg, c3, c3, "100 €", SERIF, 44, ENCRE, align="c", ancre="m")
        texte(s, x + mg, y + mg + c3, c3 * 3, c3, "Un coin.", SERIF, 84, ENCRE)
        texte(s, x + mg, y + mg + 2 * c3, c3 * 3, c3, "2 heures. Un placard, une entrée, une salle de bain.",
              SANS, 22, ENCRE, interligne=1.3)
        logo(s, "nine-logo-encre", x + mg, y + mg + 6, c3 * 1.1)
    texte(s, [C1, C2, C3][i], y + sz + 40, COL, 50, nom, SERIF, 44, ENCRE)
    texte(s, [C1, C2, C3][i], y + sz + 105, COL - 20, 100, d, SANS, 24, ENCRE, interligne=1.35)

# ══════════════════════════════════════════════════════════════════════════
# CHAPITRE 07 · LA PHOTO
# ══════════════════════════════════════════════════════════════════════════
CH = "07 · La photo"
ouverture("07", "La photo", "Des vraies photos, prises par Nine, chez de vrais clients qui ont dit oui. Rien d'autre.",
          ["Quatre principes", "L'avant / après"])

s = page(CREME)
chrome(s, CH)
titre(s, "Quatre principes.", taille=96)
for i, (t, d) in enumerate([
    ("La lumière du jour.", "Fenêtre ouverte, plafonnier éteint. Pas de flash, pas de filtre."),
    ("De face.", "L'appareil droit, à hauteur du meuble. Les lignes verticales restent verticales."),
    ("Les mains, pas les visages.", "On montre Nine au travail, pas les clients. Aucune adresse, aucun courrier lisible."),
    ("Rien de mis en scène.", "Pas de fleurs posées pour la photo, pas d'images de banque. Le vrai placard, rangé."),
]):
    x = M + i * 440
    photo(s, x, 270, 408, 420, ["Étagère, lumière du jour", "Placard, de face", "Les mains de Nine", "Tiroir, tel quel"][i])
    texte(s, x, 720, 408, 110, t, SERIF, 40, ENCRE, interligne=1.0)
    texte(s, x, 830, 400, 150, d, SANS, 23, ENCRE, interligne=1.35)

s = page(CREME)
chrome(s, CH)
titre(s, "L'avant / après.", taille=96)
chapo(s, "Le contenu qui rassure le plus. Pris du même endroit, avec la même lumière, le même cadrage : "
         "on marque l'emplacement des pieds au scotch kraft.", y=250, w=1500, taille=28)
photo(s, M, 400, 560, 480, "Avant")
photo(s, M + 572, 400, 560, 480, "Après")
etiquette(s, M + 24, 400 + 480 - 64, "avant")
etiquette(s, M + 572 + 24, 400 + 480 - 64, "après · 2 h", fond=TOMATE, coul=ENCRE)
for k, (t, d) in enumerate([
    ("Même place", "Un repère au sol, l'appareil à la même hauteur."),
    ("Même lumière", "Même heure si possible, même source."),
    ("Accord écrit", "Deux autorisations distinctes : prendre la photo, puis la publier."),
]):
    y = 400 + k * 165
    texte(s, 1320, y, 504, 50, t, SERIF, 40, ENCRE)
    texte(s, 1320, y + 56, 504, 90, d, SANS, 24, ENCRE, interligne=1.35)

# ══════════════════════════════════════════════════════════════════════════
# CHAPITRE 08 · LE TRAIT ET LA MATIÈRE
# ══════════════════════════════════════════════════════════════════════════
CH = "08 · Le trait et la matière"
ouverture("08", "Le trait et la matière", "Des pictogrammes au trait, comme sur le site actuel. Et les vraies matières du métier : le kraft, le carton, l'étiqueteuse.",
          ["Les pictogrammes", "L'étiquette et le kraft", "Le mouvement"])

s = page(CREME)
chrome(s, CH)
titre(s, "Les pictogrammes.", taille=96)
chapo(s, "Un trait encre d'épaisseur constante, des angles droits, des objets du quotidien. "
         "Le carré tomate apparaît une fois par picto, au plus.", y=250, w=1500, taille=28)
ep = 7
sz = 200
for i, nom in enumerate(["Un coin", "Une pièce", "Réserver", "La boîte", "Le tiroir", "L'étiquette"]):
    x = M + i * 292
    y = 420
    rect(s, x, y, 260, 300, BLANC)
    cx, cy = x + 30, y + 40
    if i == 0:
        rect(s, cx + 20, cy + 10, 160, 160, None, ENCRE, ep)
        rect(s, cx + 20 + 160 * 2 / 3, cy + 10, 160 / 3, 160 / 3, TOMATE)
    elif i == 1:
        rect(s, cx + 50, cy, 100, 180, None, ENCRE, ep)
        rect(s, cx + 120, cy + 90, 14, 14, ENCRE)
        ligne(s, cx + 10, cy + 180, cx + 190, cy + 180, ENCRE, ep)
    elif i == 2:
        rect(s, cx + 20, cy + 20, 160, 150, None, ENCRE, ep)
        ligne(s, cx + 20, cy + 60, cx + 180, cy + 60, ENCRE, ep)
        ligne(s, cx + 60, cy, cx + 60, cy + 36, ENCRE, ep)
        ligne(s, cx + 140, cy, cx + 140, cy + 36, ENCRE, ep)
        for r_ in range(2):
            for c_ in range(3):
                rect(s, cx + 45 + c_ * 42, cy + 82 + r_ * 42, 24, 24, TOMATE if (r_, c_) == (0, 2) else ENCRE)
    elif i == 3:
        rect(s, cx + 10, cy + 30, 180, 40, None, ENCRE, ep)
        rect(s, cx + 22, cy + 70, 156, 110, None, ENCRE, ep)
        rect(s, cx + 70, cy + 105, 60, 30, None, ENCRE, ep)
    elif i == 4:
        rect(s, cx + 10, cy + 10, 180, 170, None, ENCRE, ep)
        ligne(s, cx + 10, cy + 95, cx + 190, cy + 95, ENCRE, ep)
        ligne(s, cx + 75, cy + 52, cx + 125, cy + 52, ENCRE, ep)
        ligne(s, cx + 75, cy + 137, cx + 125, cy + 137, ENCRE, ep)
    elif i == 5:
        rect(s, cx + 10, cy + 50, 180, 90, None, ENCRE, ep, forme=MSO_SHAPE.PENTAGON)
        rect(s, cx + 150, cy + 85, 18, 18, TOMATE)
    texte(s, x, y + 230, 260, 40, nom, SANS, 24, ENCRE, gras=True, align="c")
texte(s, M, 780, 1728, 200,
      [[("Construction. ", {"gras": True}),
        ("Pictos dessinés sur un carré de 24 unités, trait de 2 unités, sans remplissage. "
         "Dans Canva, on garde ces six-là comme modèles et on dessine les suivants de la même façon, "
         "plutôt que de piocher dans les éléments tout faits.", {})]],
      SANS, 26, ENCRE, interligne=1.4)

s = page(KRAFT)
chrome(s, CH, sombre=False)
titre(s, "L'étiquette\net le kraft.", w=900, taille=96)
chapo(s, "L'étiqueteuse est l'outil de Nine. Elle devient un élément de la marque : une bande encre, "
         "des capitales espacées, et on étiquette tout, y compris les visuels.", y=380, w=780, taille=28)
for k, (t, f, c) in enumerate([
    ("à garder", ENCRE, CREME), ("à donner", ENCRE, CREME), ("à jeter", ENCRE, CREME),
    ("réserver", TOMATE, ENCRE), ("2 heures · 100 €", CREME, ENCRE), ("on range, on ne nettoie pas", ENCRE, CREME),
]):
    etiquette(s, 1000 + (k % 2) * 30, 200 + k * 110, t, fond=f, coul=c, taille=34)
texte(s, M, 760, 820, 260,
      [[("Usages. ", {"gras": True}),
        ("Titres de rubrique, prix, avant / après, stickers sur les boîtes, bandeau de story. "
         "Le texte de l'étiquette est toujours un fait, jamais une blague.", {})]],
      SANS, 26, ENCRE, interligne=1.4)
texte(s, 1000, 900, 824, 100,
      "Fond kraft : une vraie photo de papier ou de carton, de préférence à l'aplat.", SANS, 22, ENCRE, interligne=1.35)

s = page(CREME)
chrome(s, CH)
titre(s, "Le mouvement :\nça se range.", w=900, taille=96)
chapo(s, "Pour les reels, les stories et le site. Les éléments glissent et s'emboîtent dans la grille, "
         "comme des boîtes sur une étagère. Le coin tomate arrive en dernier.", y=380, w=820, taille=28)
for k in range(4):
    x = 1000 + k * 210
    rect(s, x, 220, 190, 190, None, CARTON, 2)
    for j in range(k + 1):
        if j < 3:
            rect(s, x + 10 + (j % 3) * 58, 220 + 132 - 0, 52, 48, ENCRE)
    if k == 3:
        rect(s, x + 10 + 2 * 58, 220 + 10, 52, 52, TOMATE)
    texte(s, x, 425, 190, 30, ["0 ms", "120 ms", "240 ms", "400 ms"][k], SANS, 20, CRAYON, align="c")
for i, (t, d) in enumerate([
    ("Court.", "Une animation dure entre 200 et 400 ms. Un reel « 1 tiroir, 1 transformation », 15 secondes."),
    ("Net.", "Ça glisse et ça s'arrête. Pas de rebond, pas de rotation, pas d'effet « magie »."),
    ("Le coin en dernier.", "La touche tomate apparaît quand tout est rangé : c'est la signature."),
]):
    y = 620 + i * 130
    texte(s, M, y, 760, 50, t, SERIF, 40, ENCRE)
    texte(s, M + 380, y + 4, 1300, 90, d, SANS, 24, ENCRE, interligne=1.35)

# ══════════════════════════════════════════════════════════════════════════
# CHAPITRE 09 · EN PRATIQUE
# ══════════════════════════════════════════════════════════════════════════
CH = "09 · En pratique"
ouverture("09", "En pratique", "La charte appliquée aux supports du lancement : de quoi démarrer dans Canva dès maintenant.",
          ["Carte et flyer", "Instagram", "Stories et profil", "Le site", "Tenue et matériel", "Avant de publier"])

# Carte de visite + flyer
s = page(CARTON)
chrome(s, CH)
titre(s, "Carte de visite\net flyer.", w=700, taille=88)
texte(s, M, 380, 640, 200, "Carte 85 × 55 mm, flyer A5. Papier kraft ou blanc mat non couché : on doit sentir la matière.",
      SANS, 26, ENCRE, interligne=1.4)
# carte recto
cw, chh = 510, 330
cx, cy = 820, 150
rect(s, cx, cy, cw, chh, ENCRE)
logo(s, "nine-signature-creme", cx + 105, cy + 90, 300)
# carte verso
cy2 = cy + chh + 40
rect(s, cx, cy2, cw, chh, CREME)
rect(s, cx + cw - 70, cy2, 70, 70, TOMATE)
texte(s, cx + 36, cy2 + 36, 400, 60, "Nine Marquer", SERIF, 44, ENCRE)
texte(s, cx + 36, cy2 + 100, 420, 40, "Vous réservez. Nine range.", SANS, 22, CRAYON)
texte(s, cx + 36, cy2 + 200, 440, 100, ["06 15 24 08 25", "nine.marquer@gmail.com", "Saint-Nazaire · Nantes"], SANS, 21, ENCRE,
      interligne=1.3)
# flyer A5
fx, fy, fw = 1400, 150, 424
fh = fw * 210 / 148
rect(s, fx, fy, fw, fh, CREME, ENCRE, 1)
rect(s, fx + fw * 2 / 3, fy, fw / 3, fw / 3, TOMATE)
logo(s, "nine-logo-encre", fx + 28, fy + 28, 150)
texte(s, fx + 28, fy + 150, fw - 56, 190, "Vous n'aimez pas ranger. Nine, si.", SERIF, 46, ENCRE, interligne=0.95)
texte(s, fx + 28, fy + 340, fw - 56, 60, "Je viens, je range, c'est réglé. Réservation en trois clics, prix affiché.",
      SANS, 17, ENCRE, interligne=1.3)
ligne(s, fx + 28, fy + 410, fx + fw - 28, fy + 410, ENCRE, 1.5)
texte(s, fx + 28, fy + 422, (fw - 56) / 2, 80, [[("Un coin", {"police": SERIF, "taille": 30})], "2 h · 100 €"], SANS, 17, ENCRE)
texte(s, fx + 28 + (fw - 56) / 2, fy + 422, (fw - 56) / 2, 80, [[("Une pièce", {"police": SERIF, "taille": 30})], "½ journée · 200 €"],
      SANS, 17, ENCRE)
etiquette(s, fx + 28, fy + fh - 60, "réserver · 06 15 24 08 25", taille=16, fond=ENCRE)
texte(s, M, 980, 700, 40, "Carte recto, carte verso, flyer A5 recto.", SANS, 20, CRAYON)

# Instagram
s = page(CREME)
chrome(s, CH)
titre(s, "Instagram : trois modèles.", taille=88)
chapo(s, "Trois publications par semaine, toujours sur ces trois modèles. Le fil se reconnaît de loin, comme une étagère bien rangée.",
      y=230, w=1600, taille=26)
sz = 520
for i in range(3):
    x = [C1, C2, C3][i] + (COL - sz) / 2
    y = 360
    mg = sz / 12
    c3 = (sz - 2 * mg) / 3
    if i == 0:
        rect(s, x, y, sz, sz, ENCRE)
        rect(s, x + sz - c3 * 0.5 - mg, y + mg, c3 * 0.5, c3 * 0.5, TOMATE)
        texte(s, x + mg, y + mg + c3 * 0.9, sz - 2 * mg, c3 * 2, "Ranger, ça se délègue.", SERIF, 74, CREME, interligne=0.95)
        logo(s, "nine-logo-creme", x + mg, y + sz - mg - 50, 130)
    elif i == 1:
        photo(s, x, y, sz / 2 - 2, sz, "Avant")
        photo(s, x + sz / 2 + 2, y, sz / 2 - 2, sz, "Après")
        etiquette(s, x + mg, y + mg, "tiroir de cuisine · 40 min", taille=16)
        etiquette(s, x + sz / 2 + mg, y + sz - mg - 32, "après", taille=16, fond=TOMATE, coul=ENCRE)
    else:
        rect(s, x, y, sz, sz, CREME, ENCRE, 2)
        rect(s, x + sz - mg - c3, y + mg, c3, c3, TOMATE)
        texte(s, x + sz - mg - c3, y + mg, c3, c3, [[("100 €", {})], [("2 heures", {"police": SANS, "taille": 20, "gras": True})]],
              SERIF, 44, ENCRE, align="c", ancre="m", interligne=1.0)
        logo(s, "nine-logo-encre", x + mg, y + mg + 10, 140)
        texte(s, x + mg, y + mg + c3 + 10, sz - 2 * mg, c3, "Un coin.", SERIF, 100, ENCRE)
        texte(s, x + mg, y + mg + 2 * c3, sz - 2 * mg, c3, "Un placard, une entrée, une salle de bain. Vous réservez, Nine range.",
              SANS, 22, ENCRE, interligne=1.3)
    texte(s, [C1, C2, C3][i], y + sz + 30, COL, 50, ["La phrase", "L'avant / après", "L'offre"][i], SERIF, 40, ENCRE)
    texte(s, [C1, C2, C3][i], y + sz + 90, COL, 80,
          ["Une phrase de la voix, rien d'autre.", "Le contenu qui rassure. Durée réelle sur l'étiquette.",
           "Une fois par semaine au plus."][i], SANS, 22, ENCRE)

# Stories + profil
s = page(CREME)
chrome(s, CH)
titre(s, "Stories, reels\net profil.", w=700, taille=88)
texte(s, M, 380, 640, 300,
      "Stories les jours d'intervention : le kit dans le coffre, les mains au travail, l'étiquette posée. "
      "Couverture de reel : une phrase et une étiquette, dans la même grille.",
      SANS, 26, ENCRE, interligne=1.4)
# profil
image(s, "nine-symbole", M, 720, 180, 180)
texte(s, M + 210, 740, 500, 160,
      [[("Photo de profil", {"gras": True})], "Le symbole, partout. Il reste lisible en tout petit, là où le mot ne l'est pas."],
      SANS, 22, ENCRE, interligne=1.35)
sw = 320
shh = sw * 16 / 9
for i in range(3):
    x = 820 + i * (sw + 32)
    y = 130
    if i == 0:
        photo(s, x, y, sw, shh, "Le kit dans le coffre")
        etiquette(s, x + 24, y + 40, "mardi · Nantes", taille=18)
        etiquette(s, x + 24, y + shh - 90, "en route", taille=18, fond=TOMATE, coul=ENCRE)
    elif i == 1:
        rect(s, x, y, sw, shh, TOMATE)
        texte(s, x + 28, y + 110, sw - 56, 260, "Le placard ne va pas se ranger tout seul.", SERIF, 40, ENCRE, interligne=0.95)
        texte(s, x + 28, y + 380, sw - 56, 60, "On a vérifié.", SERIF, 40, CREME)
        logo(s, "nine-logo-sur-tomate", x + 28, y + shh - 80, 110)
    else:
        rect(s, x, y, sw, shh, ENCRE)
        photo(s, x, y + 120, sw, sw, "Tiroir, vue du dessus", fond=CRAYON)
        texte(s, x + 28, y + 120 + sw + 24, sw - 56, 110, "1 tiroir,\n1 transformation.", SERIF, 34, CREME, interligne=0.95)
        etiquette(s, x + 28, y + 40, "15 s", taille=18, fond=CREME, coul=ENCRE)
    texte(s, x, y + shh + 16, sw, 30, ["Story d'intervention", "Story phrase", "Couverture de reel"][i], SANS, 20, CRAYON)

# Le site
s = page(CREME)
chrome(s, CH)
titre(s, "Le site.", taille=96)
chapo(s, "On garde la voix et le fond crème du site actuel. On change l'accroche, les offres et la réservation.",
      y=240, w=660, taille=26)
bx, by, bw, bh = 820, 130, 1004, 860
rect(s, bx, by, bw, bh, BLANC, ENCRE, 2)
rect(s, bx, by, bw, 40, CARTON)
for k in range(3):
    rect(s, bx + 18 + k * 22, by + 14, 12, 12, ENCRE if k else TOMATE)
logo(s, "nine-logo-encre", bx + 40, by + 70, 110)
texte(s, bx + bw - 440, by + 76, 300, 30, "Les offres   Nine   Questions", SANS, 18, ENCRE)
etiquette(s, bx + bw - 40, by + 70, "réserver", taille=18, fond=TOMATE, coul=ENCRE, ancre_droite=True)
texte(s, bx + 40, by + 170, 640, 300, "Vous n'aimez pas ranger. Nine, si.", SERIF, 84, ENCRE, interligne=0.92)
texte(s, bx + 40, by + 440, 560, 80, "Vous réservez un créneau, je viens, je range. Pas de devis, pas de diagnostic.",
      SANS, 22, ENCRE, interligne=1.35)
photo(s, bx + 700, by + 170, 264, 330, "Nine au travail")
for k, (n, d) in enumerate([("Un coin", "2 h · 100 €"), ("Une pièce", "½ journée · 200 €")]):
    x = bx + 40 + k * 470
    rect(s, x, by + 560, 450, 250, ENCRE if k == 0 else None, None if k == 0 else ENCRE, 2)
    c = CREME if k == 0 else ENCRE
    texte(s, x + 30, by + 590, 380, 70, n, SERIF, 56, c)
    texte(s, x + 30, by + 670, 380, 40, d, SANS, 24, c, gras=True)
    texte(s, x + 30, by + 750, 380, 40, "RÉSERVER →", SANS, 18, TOMATE if k == 0 else ENCRE, gras=True, espacement=3)
for k, t in enumerate(["Accroche : « Vous n'aimez pas ranger. Nine, si. »", "Deux offres en vitrine, prix affiché",
                       "Réservation en trois clics, dès 48 h", "Plus de diagnostic, de caméra ni de prix barrés"]):
    coche(s, M, 420 + k * 90, 28, TOMATE if k == 0 else ENCRE)
    texte(s, M + 48, 420 + k * 90, 640, 80, t, SANS, 24, ENCRE, interligne=1.3)

# Tenue et matériel
s = page(CREME)
chrome(s, CH)
titre(s, "Tenue et matériel.", taille=96)
chapo(s, "Nine entre chez les gens : la marque doit la rendre identifiable et rassurante, pas déguisée.",
      y=250, w=1500, taille=28)
# t-shirt (forme simple)
tx, ty = M + 40, 380
rect(s, tx + 90, ty, 300, 420, ENCRE)
rect(s, tx, ty, 110, 150, ENCRE)
rect(s, tx + 370, ty, 110, 150, ENCRE)
rect(s, tx + 190, ty, 100, 30, CREME)
image(s, "nine-symbole-creme", tx + 290, ty + 70, 50, 50)
texte(s, M, 830, 560, 120, [[("Le haut. ", {"gras": True}), ("Encre, symbole crème brodé côté cœur. Deux exemplaires.", {})]],
      SANS, 24, ENCRE, interligne=1.35)
# boîte kraft + sticker
bx, by = 760, 420
rect(s, bx, by, 420, 80, KRAFT)
rect(s, bx + 20, by + 80, 380, 300, KRAFT)
ligne(s, bx + 20, by + 80, bx + 400, by + 80, ENCRE, 2)
rect(s, bx + 150, by + 170, 120, 120, CREME)
image(s, "nine-symbole", bx + 170, by + 190, 80, 80)
texte(s, 700, 830, 520, 120, [[("Le sticker. ", {"gras": True}), ("Le symbole sur fond crème, 5 cm, sur la caisse de matériel.", {})]],
      SANS, 24, ENCRE, interligne=1.35)
# étiquettes
for k, t in enumerate(["à garder", "à donner", "à jeter"]):
    etiquette(s, 1340, 440 + k * 90, t, taille=34, fond=CREME if k == 2 else ENCRE, coul=ENCRE if k == 2 else CREME)
texte(s, 1340, 830, 484, 120, [[("Les sacs de tri. ", {"gras": True}), ("Étiquetés devant le client : c'est déjà la marque.", {})]],
      SANS, 24, ENCRE, interligne=1.35)

# Avant de publier
s = page(ENCRE)
chrome(s, CH, sombre=True)
texte(s, M, 120, 760, 300, "Avant de publier :\nneuf questions.", SERIF, 88, CREME, interligne=0.95)
texte(s, M, 360, 640, 200, "Neuf oui, on publie. Un non, on range encore un peu.", SANS, 28, KRAFT, interligne=1.35)
questions = [
    "Une seule touche de tomate ?", "Le titre est une phrase, avec un point ?", "Tout tient dans la grille ?",
    "Aucun mot de la liste « on évite » ?", "Lisible en une seconde ?", "Le logo a de l'air autour ?",
    "Ni visage ni adresse de client ?", "Lumière du jour, de face ?", "Les prix sont ceux en vigueur ?",
]
gx, gy, cs = 880, 110, 286
cs_w = (W - M - gx - 24) / 3
for i, q in enumerate(questions):
    r, c = divmod(i, 3)
    x, y = gx + c * (cs_w + 12), gy + r * (cs + 12)
    coin = (r, c) == (0, 2)
    rect(s, x, y, cs_w, cs, TOMATE if coin else None, None if coin else CREME, 1.5)
    texte(s, x + 24, y + 20, 80, 60, f"{i + 1}", SERIF, 52, ENCRE if coin else KRAFT)
    texte(s, x + 24, y + cs - 140, cs_w - 48, 120, q, SANS, 25, ENCRE if coin else CREME, gras=True, interligne=1.25,
          ancre="b")

# Ce qui reste à décider
s = page(CREME)
chrome(s, "Pour finir")
titre(s, "Ce qui reste à décider.", taille=96)
chapo(s, "Cette version 1 tranche le système. Quelques choix restent à confirmer avec Nine, sur pièce.",
      y=250, w=1500, taille=28)
reste = [
    ("La tomate, sur papier.", "Imprimer un échantillon (carte, sticker) avant de valider la couleur : un écran ment toujours un peu."),
    ("Le nom de domaine.", "nine.fr et nine.com sont pris. Pistes déjà notées : nine44.fr, nine-services.fr. À choisir avant les cartes."),
    ("Les vraies photos.", "Tous les emplacements « PHOTO » de ce document attendent les prestations tests de décembre."),
    ("Le site actuel.", "Sa couleur d'accent et ses polices sont à aligner sur cette charte lors de la refonte."),
    ("Les mentions fiscales.", "Le prix après crédit d'impôt apparaît partout dès le récépissé NOVA reçu, pas avant."),
    ("La version 2.", "Après le bilan de janvier : ce qui a marché sur les réseaux devient la règle."),
]
for i, (t, d) in enumerate(reste):
    r, c = divmod(i, 3)
    x, y = [C1, C2, C3][c], 380 + r * 320
    ligne(s, x, y, x + COL, y, ENCRE, 3)
    texte(s, x, y + 24, COL, 60, t, SERIF, 40, ENCRE)
    texte(s, x, y + 90, COL - 20, 200, d, SANS, 24, ENCRE, interligne=1.4)

# Dos de couverture
s = page(ENCRE)
rect(s, W - 360, 0, 360, 360, TOMATE)
image(s, "nine-signature-creme", M, 340, w=700)
texte(s, M, 830, 1200, 100, "Vous réservez. Nine range.", SERIF, 64, CREME)
texte(s, M, 960, 1200, 40, "Charte de marque · version 1 · octobre 2026", SANS, 22, CRAYON, gras=True, espacement=2)

prs.save(SORTIE)
print(f"{len(pages)} pages → {SORTIE}")
